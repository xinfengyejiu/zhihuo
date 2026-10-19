"""知识库(M2):上传解析、分块、Embedding、检索测试、挂载到书。"""

import html as html_mod
import io
import json
import re
import zipfile

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from sqlalchemy.orm import Session

from ..ai.gateway import embed_texts
from ..ai.retrieval import chunk_text, retrieve
from ..database import get_db
from ..models import KbChunk, KbDocument, KnowledgeBase, NovelKbLink, ResourceMount, User
from ..routers.novels import get_owned_novel
from ..schemas import KbCreateIn, KbDocumentOut, KbMountIn, KbOut, KbSearchHit, KbSearchIn
from ..security import get_current_user

router = APIRouter(prefix="/kb", tags=["kb"])
mount_router = APIRouter(prefix="/novels/{novel_id}/kbs", tags=["kb"])

MAX_FILE_SIZE = 10 * 1024 * 1024  # 10MB
MAX_FOLDER_BYTES = 50 * 1024 * 1024  # 文件夹合计 50MB
MAX_FOLDER_FILES = 100
MAX_CHUNKS = 3000
ALLOWED_SUFFIXES = ("txt", "md", "markdown", "pdf", "docx", "epub")


def _normalize_rel_path(path: str) -> str:
    p = (path or "").replace("\\", "/").lstrip("/")
    parts = [x for x in p.split("/") if x and x != "." and x != ".."]
    return "/".join(parts)[:255] or "untitled.txt"


def _is_allowed_filename(filename: str) -> bool:
    name = (filename or "").lower().rsplit("/", 1)[-1]
    if name.startswith(".") or name in ("__macosx",):
        return False
    if "/__macosx/" in f"/{(filename or '').lower()}/":
        return False
    suffix = name.rsplit(".", 1)[-1] if "." in name else ""
    return suffix in ALLOWED_SUFFIXES


@router.get("", response_model=list[KbOut])
def list_kbs(
    scope: str | None = None,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """列出知识库。scope=novel|media|drama 时只返回该产品或 shared 标签。"""
    kbs = db.query(KnowledgeBase).filter_by(user_id=user.id).order_by(KnowledgeBase.id).all()
    if scope:
        key = scope.strip().lower()
        filtered = []
        for kb in kbs:
            tags = kb.scope_tags or ["shared"]
            if key in tags or "shared" in tags:
                filtered.append(kb)
        kbs = filtered
    out = []
    for kb in kbs:
        docs = db.query(KbDocument).filter_by(kb_id=kb.id).all()
        doc_ids = [d.id for d in docs]
        chunk_count = db.query(KbChunk).filter(KbChunk.kb_id == kb.id).count() if doc_ids else 0
        out.append(
            KbOut(
                id=kb.id,
                name=kb.name,
                description=kb.description,
                scope_tags=kb.scope_tags or [],
                document_count=len(docs),
                chunk_count=chunk_count,
            )
        )
    return out


@router.post("", response_model=KbOut)
def create_kb(body: KbCreateIn, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    tags = body.scope_tags or ["shared"]
    kb = KnowledgeBase(user_id=user.id, name=body.name, description=body.description, scope_tags=tags)
    db.add(kb)
    db.commit()
    db.refresh(kb)
    return KbOut(id=kb.id, name=kb.name, description=kb.description, scope_tags=kb.scope_tags or [])


@router.delete("/{kb_id}")
def delete_kb(kb_id: int, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    kb = db.query(KnowledgeBase).filter_by(id=kb_id, user_id=user.id).first()
    if kb is None:
        raise HTTPException(404, "知识库不存在")
    db.delete(kb)
    db.commit()
    return {"ok": True}


def _extract_text(filename: str, data: bytes) -> str:
    suffix = filename.lower().rsplit(".", 1)[-1]
    if suffix in ("txt", "md", "markdown"):
        for enc in ("utf-8", "gb18030", "utf-16"):
            try:
                return data.decode(enc)
            except UnicodeDecodeError:
                continue
        return data.decode("utf-8", errors="replace")
    if suffix == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if suffix == "docx":
        import docx

        doc = docx.Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append("\t".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    if suffix == "epub":
        parts = []
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            for name in zf.namelist():
                if name.lower().endswith((".xhtml", ".html", ".htm")):
                    raw = zf.read(name).decode("utf-8", errors="replace")
                    text = re.sub(r"<(script|style)[^>]*>.*?</\1>", "", raw, flags=re.S | re.I)
                    text = re.sub(r"<br\s*/?>", "\n", text, flags=re.I)
                    text = re.sub(r"</(p|div|h[1-6]|li)>", "\n", text, flags=re.I)
                    text = re.sub(r"<[^>]+>", "", text)
                    text = html_mod.unescape(text)
                    parts.append(text)
        return "\n".join(parts)
    raise HTTPException(400, f"暂不支持的文件类型:.{suffix}(支持 TXT/MD/PDF/DOCX/EPUB)")


async def _embed_chunks(db: Session, user: User, chunks: list[str]) -> list[list | None]:
    """分批向量化;Embedding 未配置时降级为仅关键词可用(向量留空)。"""
    from ..ai.gateway import resolve_route, AiNotConfigured

    try:
        resolve_route(db, user.id, "embedding")
    except AiNotConfigured:
        return [None] * len(chunks)
    vectors: list[list | None] = []
    batch = 32
    for i in range(0, len(chunks), batch):
        vs = await embed_texts(db, user.id, chunks[i : i + batch])
        vectors.extend(vs)
    return vectors


@router.get("/{kb_id}/documents", response_model=list[KbDocumentOut])
def list_documents(kb_id: int, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    kb = db.query(KnowledgeBase).filter_by(id=kb_id, user_id=user.id).first()
    if kb is None:
        raise HTTPException(404, "知识库不存在")
    return db.query(KbDocument).filter_by(kb_id=kb_id).order_by(KbDocument.id.desc()).all()


async def _ingest_one(
    db: Session,
    user: User,
    kb_id: int,
    filename: str,
    data: bytes,
    *,
    raise_on_error: bool = True,
) -> KbDocument:
    """解析单文件并分块入库。raise_on_error=False 时失败文档也提交,便于批量导入。"""
    display = _normalize_rel_path(filename)
    if len(data) > MAX_FILE_SIZE:
        if raise_on_error:
            raise HTTPException(400, f"{display}:文件超过 10MB 限制")
        doc = KbDocument(kb_id=kb_id, filename=display, size=len(data), status="failed", error="超过 10MB")
        db.add(doc)
        db.flush()
        return doc
    if not _is_allowed_filename(display):
        if raise_on_error:
            raise HTTPException(400, f"暂不支持的文件类型:{display}")
        doc = KbDocument(
            kb_id=kb_id,
            filename=display,
            size=len(data),
            status="failed",
            error="不支持的类型(仅 TXT/MD/PDF/DOCX/EPUB)",
        )
        db.add(doc)
        db.flush()
        return doc

    doc = KbDocument(kb_id=kb_id, filename=display, size=len(data), status="failed")
    db.add(doc)
    db.flush()
    try:
        text = _extract_text(display, data)
        chunks = chunk_text(text)
        if not chunks:
            raise ValueError("未能从文件中提取到文本")
        if len(chunks) > MAX_CHUNKS:
            chunks = chunks[:MAX_CHUNKS]
        vectors = await _embed_chunks(db, user, chunks)
        for idx, (content, vec) in enumerate(zip(chunks, vectors)):
            db.add(KbChunk(kb_id=kb_id, document_id=doc.id, idx=idx, content=content, embedding=vec))
        doc.chunk_count = len(chunks)
        doc.status = "ready"
        if all(v is None for v in vectors):
            doc.error = "Embedding 未配置,已分块入库但暂不可向量检索"
    except HTTPException as e:
        detail = e.detail if isinstance(e.detail, str) else "文件解析失败"
        doc.error = detail[:500]
        if raise_on_error:
            db.commit()
            raise
    except Exception as e:  # noqa: BLE001
        doc.error = f"处理失败:{e}"[:500]
        if raise_on_error:
            db.commit()
            raise HTTPException(400, doc.error)
    return doc


@router.post("/{kb_id}/upload", response_model=KbDocumentOut)
async def upload_document(
    kb_id: int,
    file: UploadFile,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    kb = db.query(KnowledgeBase).filter_by(id=kb_id, user_id=user.id).first()
    if kb is None:
        raise HTTPException(404, "知识库不存在")
    data = await file.read()
    doc = await _ingest_one(db, user, kb_id, file.filename or "untitled.txt", data, raise_on_error=True)
    db.commit()
    db.refresh(doc)
    return doc


@router.post("/{kb_id}/import-folder")
async def import_folder(
    kb_id: int,
    files: list[UploadFile] = File(...),
    paths: str = Form("[]"),
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """浏览器选择文件夹上传:paths 为 JSON 相对路径数组,与 files 一一对应。"""
    kb = db.query(KnowledgeBase).filter_by(id=kb_id, user_id=user.id).first()
    if kb is None:
        raise HTTPException(404, "知识库不存在")
    try:
        rel_paths = json.loads(paths)
    except json.JSONDecodeError as e:
        raise HTTPException(400, "paths 必须是 JSON 数组") from e
    if not isinstance(rel_paths, list) or len(rel_paths) != len(files):
        raise HTTPException(400, "files 与 paths 数量不一致")
    if len(files) > MAX_FOLDER_FILES:
        raise HTTPException(400, f"单次最多导入 {MAX_FOLDER_FILES} 个文件")

    items: list[tuple[str, bytes]] = []
    skipped: list[str] = []
    total = 0
    for upload, rel in zip(files, rel_paths):
        name = _normalize_rel_path(str(rel) or (upload.filename or ""))
        if not _is_allowed_filename(name):
            skipped.append(name or (upload.filename or "?"))
            continue
        raw = await upload.read()
        total += len(raw)
        if total > MAX_FOLDER_BYTES:
            raise HTTPException(400, f"文件夹过大(上限 {MAX_FOLDER_BYTES // (1024 * 1024)}MB)")
        items.append((name, raw))

    if not items:
        raise HTTPException(
            400,
            "文件夹内没有可导入文件(支持 TXT/MD/PDF/DOCX/EPUB)"
            + (f";已跳过 {len(skipped)} 个" if skipped else ""),
        )

    created: list[dict] = []
    failed: list[dict] = []
    for name, raw in items:
        doc = await _ingest_one(db, user, kb_id, name, raw, raise_on_error=False)
        if doc.status == "ready":
            created.append({"id": doc.id, "filename": doc.filename, "chunk_count": doc.chunk_count})
        else:
            failed.append({"filename": doc.filename, "error": doc.error})
    db.commit()
    return {
        "ok": True,
        "created": len(created),
        "failed": len(failed),
        "skipped": len(skipped),
        "items": created,
        "errors": failed[:20],
        "skipped_names": skipped[:30],
    }


@router.post("/{kb_id}/import-zip")
async def import_zip(
    kb_id: int,
    file: UploadFile = File(...),
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """上传 zip 包,解压后按相对路径批量入库(等同文件夹导入)。"""
    kb = db.query(KnowledgeBase).filter_by(id=kb_id, user_id=user.id).first()
    if kb is None:
        raise HTTPException(404, "知识库不存在")
    data = await file.read()
    if len(data) > MAX_FOLDER_BYTES:
        raise HTTPException(400, f"zip 过大(上限 {MAX_FOLDER_BYTES // (1024 * 1024)}MB)")
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as e:
        raise HTTPException(400, "不是有效的 zip 文件") from e

    items: list[tuple[str, bytes]] = []
    skipped: list[str] = []
    total = 0
    with zf:
        for info in zf.infolist():
            if info.is_dir():
                continue
            name = _normalize_rel_path(info.filename)
            if not _is_allowed_filename(name):
                skipped.append(name)
                continue
            if len(items) >= MAX_FOLDER_FILES:
                raise HTTPException(400, f"单次最多导入 {MAX_FOLDER_FILES} 个文件")
            raw = zf.read(info)
            total += len(raw)
            if total > MAX_FOLDER_BYTES:
                raise HTTPException(400, f"解压后过大(上限 {MAX_FOLDER_BYTES // (1024 * 1024)}MB)")
            items.append((name, raw))

    if not items:
        raise HTTPException(400, "zip 内没有可导入文件(支持 TXT/MD/PDF/DOCX/EPUB)")

    created: list[dict] = []
    failed: list[dict] = []
    for name, raw in items:
        doc = await _ingest_one(db, user, kb_id, name, raw, raise_on_error=False)
        if doc.status == "ready":
            created.append({"id": doc.id, "filename": doc.filename, "chunk_count": doc.chunk_count})
        else:
            failed.append({"filename": doc.filename, "error": doc.error})
    db.commit()
    return {
        "ok": True,
        "created": len(created),
        "failed": len(failed),
        "skipped": len(skipped),
        "items": created,
        "errors": failed[:20],
        "skipped_names": skipped[:30],
    }


@router.delete("/{kb_id}/documents/{doc_id}")
def delete_document(kb_id: int, doc_id: int, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    kb = db.query(KnowledgeBase).filter_by(id=kb_id, user_id=user.id).first()
    if kb is None:
        raise HTTPException(404, "知识库不存在")
    doc = db.query(KbDocument).filter_by(id=doc_id, kb_id=kb_id).first()
    if doc is None:
        raise HTTPException(404, "文档不存在")
    db.delete(doc)
    db.commit()
    return {"ok": True}


@router.post("/{kb_id}/search", response_model=list[KbSearchHit])
async def search_kb(
    kb_id: int, body: KbSearchIn, user: User = Depends(get_current_user), db: Session = Depends(get_db)
):
    kb = db.query(KnowledgeBase).filter_by(id=kb_id, user_id=user.id).first()
    if kb is None:
        raise HTTPException(404, "知识库不存在")
    # 直接在该库内检索(不经过书挂载)
    chunks = (
        db.query(KbChunk)
        .filter(KbChunk.kb_id == kb_id, KbChunk.embedding.isnot(None))
        .limit(5000)
        .all()
    )
    if not chunks:
        if db.query(KbChunk).filter(KbChunk.kb_id == kb_id).count() > 0:
            raise HTTPException(400, "该库尚未向量化:请先在 个人中心 → 用途路由 配置 Embedding 模型,再重新上传文档")
        return []
    qvec = (await embed_texts(db, user.id, [body.query[:1000]]))[0]
    from ..ai.retrieval import _cosine

    scored = sorted(((c, _cosine(qvec, c.embedding)) for c in chunks), key=lambda p: p[1], reverse=True)
    return [
        KbSearchHit(content=c.content, document_id=c.document_id, kb_id=c.kb_id, score=round(s, 4))
        for c, s in scored[: max(1, min(body.top_k, 20))]
        if s > 0.05
    ]


# ---------- 挂载到书 ----------
@mount_router.get("")
def get_mounted(novel_id: int, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    get_owned_novel(novel_id, user, db)
    links = db.query(NovelKbLink).filter_by(novel_id=novel_id).all()
    return {"kb_ids": [l.kb_id for l in links]}


@mount_router.put("")
def set_mounted(
    novel_id: int, body: KbMountIn, user: User = Depends(get_current_user), db: Session = Depends(get_db)
):
    get_owned_novel(novel_id, user, db)
    valid_ids = {
        kb.id
        for kb in db.query(KnowledgeBase).filter(
            KnowledgeBase.user_id == user.id, KnowledgeBase.id.in_(body.kb_ids or [])
        ).all()
    }
    db.query(NovelKbLink).filter_by(novel_id=novel_id).delete()
    for kb_id in sorted(valid_ids):
        db.add(NovelKbLink(novel_id=novel_id, kb_id=kb_id))
    # 同步 ResourceMount
    db.query(ResourceMount).filter_by(
        user_id=user.id,
        product_key="novel",
        workspace_type="novel",
        workspace_id=novel_id,
        resource_type="kb",
    ).delete(synchronize_session=False)
    for kb_id in sorted(valid_ids):
        db.add(
            ResourceMount(
                user_id=user.id,
                product_key="novel",
                workspace_type="novel",
                workspace_id=novel_id,
                resource_type="kb",
                resource_id=kb_id,
                stage="",
            )
        )
    db.commit()
    return {"ok": True, "kb_ids": sorted(valid_ids)}


@router.put("/{kb_id}/scope", response_model=KbOut)
def update_kb_scope(
    kb_id: int,
    body: KbCreateIn,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    kb = db.query(KnowledgeBase).filter_by(id=kb_id, user_id=user.id).first()
    if kb is None:
        raise HTTPException(404, "知识库不存在")
    if body.name:
        kb.name = body.name
    if body.description is not None:
        kb.description = body.description
    kb.scope_tags = body.scope_tags or ["shared"]
    db.commit()
    db.refresh(kb)
    docs = db.query(KbDocument).filter_by(kb_id=kb.id).all()
    chunk_count = db.query(KbChunk).filter(KbChunk.kb_id == kb.id).count()
    return KbOut(
        id=kb.id,
        name=kb.name,
        description=kb.description,
        scope_tags=kb.scope_tags or [],
        document_count=len(docs),
        chunk_count=chunk_count,
    )
